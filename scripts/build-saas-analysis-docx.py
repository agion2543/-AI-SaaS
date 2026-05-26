from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Administrator\Documents\Codex\2026-04-27\go-web-gin-health")
OUT = ROOT / "outputs" / "本地生活商家AI运营SaaS-方案梳理与优化路线图.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, color="FFFFFF")
        set_cell_shading(hdr[i], "2563EB")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if i == 0:
                set_cell_shading(cells[i], "F8FBFF")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(15, 23, 42)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        run.bold = True
        rest = text[len(bold_prefix):]
        p.add_run(rest)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(51, 65, 85)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.75)
section.right_margin = Cm(1.75)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("本地生活商家 AI 运营 SaaS\n方案梳理与优化路线图")
run.bold = True
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(15, 23, 42)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("面向本地商家的门店、商品、扫码点单、订单财务、优惠活动与 AI 经营分析一体化平台")
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(100, 116, 139)

add_heading(doc, "一、整体定位", 1)
add_para(doc, "本项目不是单纯的点单系统，也不是普通会员后台，而是面向本地生活商家的 AI 运营 SaaS。平台方提供商家入驻、订阅、财务、结算、风控与运营管理能力；商家端负责门店、商品、订单、优惠活动、财务对账和 AI 经营分析；顾客端通过扫码进入门店，完成商品浏览、购物车、下单支付和订单状态追踪。")
add_para(doc, "短期策略：优先把交易闭环、财务口径、订单售后和商家经营看板做稳，再整理生产部署包。支付合规上，当前更适合使用沙箱/模拟支付验证流程，真实收款前需要办理合法主体资质。", "短期策略：")

add_heading(doc, "二、三端功能地图", 1)
add_table(doc, ["端口", "核心角色", "当前能力", "下一步重点"], [
    ["平台端", "平台运营方", "商家管理、订阅开通、订单财务、退款入口、人工结算、安全检查、审计日志。", "强化结算、风险预警、商家详情、生产配置修复向导。"],
    ["商家端", "门店/企业客户", "注册登录、订阅、门店、商品、订单、财务、收款设置、优惠活动、AI 分析。", "经营看板接真实订单、商品分类、图片上传、活动抵扣、订单售后增强。"],
    ["顾客端", "扫码点单消费者", "门店页、商品展示、购物车、创建订单、订单状态页雏形。", "支付失败重试、订单状态体验、评价、会员/优惠券领取。"],
], [2.2, 3, 6.3, 5.6])

add_heading(doc, "三、阶段路线图", 1)
add_table(doc, ["阶段", "优先级", "任务", "验收标准"], [
    ["第 1 阶段", "P0", "财务导出与对账稳定，所有 CSV 改为 XLSX，平台/商家财务口径统一。", "WPS/Excel 正常打开，金额口径不混入订阅订单。"],
    ["第 1 阶段", "P0", "顾客订单状态页与支付失败重试。", "顾客能看到待支付、待接单、已接单、已完成、已退款等状态。"],
    ["第 2 阶段", "P1", "商家经营看板接真实订单数据。", "展示今日订单、实收、退款、客单价、热销和低动销商品。"],
    ["第 2 阶段", "P1", "商品分类、图片上传、上下架排序，优惠券/满减接入下单金额。", "顾客下单金额能被活动真实抵扣。"],
    ["第 3 阶段", "P1", "AI 建议可操作化。", "AI 建议能一键转为优惠活动、召回券和满减策略。"],
    ["第 4 阶段", "P2", "生产部署包。", "形成 .env、Nginx、HTTPS、备份、迁移和启动脚本说明。"],
], [2.4, 2, 8.2, 5])

add_heading(doc, "四、差异化分析", 1)
add_table(doc, ["维度", "普通系统", "本项目差异化", "可形成壁垒"], [
    ["目标", "解决点单或收银。", "围绕本地商家经营增长，连接交易、顾客、活动和 AI 建议。", "行业场景理解与运营模板。"],
    ["数据", "记录订单流水。", "沉淀顾客画像、商品销量、活动效果、退款和复购数据。", "越用越能生成更贴近经营的建议。"],
    ["AI", "多为客服或文案。", "AI 建议能转为可执行活动，例如新客券、复购券、沉睡召回券。", "从建议到执行的闭环。"],
    ["商业", "一次性软件或基础订阅。", "SaaS 订阅 + 人工结算 + 后续服务商/分账扩展。", "先轻量验证，再逐步合规扩展。"],
], [2.2, 4.8, 6.2, 4.2])

add_heading(doc, "五、实践成绩", 1)
add_para(doc, "目前项目已具备平台端、商家端、顾客端的基础框架，并形成了从商家订阅、商品管理、顾客扫码点单、订单接单、退款留痕、平台财务和人工结算的雏形。")
add_para(doc, "已经完成的重要能力包括：多角色体系、商家订阅限制、顾客扫码商品展示、购物车下单、商家订单管理、商家财务、平台订单财务、XLSX 导出、安全检查面板、生产配置修复向导。")

add_heading(doc, "六、商业与合规建议", 1)
add_para(doc, "现阶段如果没有营业执照，不建议对真实用户开放正式线上收款。支付宝/微信正式商家收款通常需要经营主体资质。短期建议继续使用沙箱或模拟支付验证订单流程；如果确实要小额测试，可以采用人工线下收款并后台留痕，但这不能替代正式支付回调。")
add_para(doc, "更稳的商业路径是：前期平台统一记录订单与财务，人工结算给商家；业务验证后办理个体工商户或公司资质，开通正式商家收款产品；成熟后再考虑官方服务商、子商户或分账能力。")

add_heading(doc, "七、下一步执行建议", 1)
add_table(doc, ["排序", "任务", "原因"], [
    ["1", "顾客订单状态页 + 支付失败重试", "补齐顾客交易闭环，让顾客知道订单现在在哪里。"],
    ["2", "商家经营看板接真实订单", "提升商家端价值，让商家打开后能看到经营结果。"],
    ["3", "优惠券/满减接入下单金额", "让营销活动从展示变成真实可用。"],
    ["4", "商品分类和图片上传", "提升顾客端点单体验和商家商品管理效率。"],
    ["5", "平台风险与结算深化", "让平台端更接近真实运营后台。"],
], [1.5, 7, 7])

footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("本文件由当前 SaaS 项目阶段性梳理生成，用于开发规划、方案复盘和沟通演示。")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
